"""Reports: aggregation with filters, the review gate, exports (Excel, PDF, CSV, payroll) and delivery."""
from __future__ import annotations

import csv
import io
import logging
from collections import defaultdict
from datetime import datetime
from zoneinfo import ZoneInfo

from . import config, db, delivery, events, settings, store
from .pipeline import STATES

log = logging.getLogger("workwatch.reports")
COLS = [("present", "Present"), ("working", "Working"), ("idle", "Idle"), ("phone", "Talking on phone"),
        ("break", "Break"), ("exclusion", "Exclusion")]


class GateError(Exception):
    """Raised when a report is requested while footage is unannotated or review items are pending."""

    def __init__(self, message, blockers):
        super().__init__(message)
        self.blockers = blockers


def hms(s: float) -> str:
    s = int(round(s or 0))
    return f"{s // 3600}:{s % 3600 // 60:02d}:{s % 60:02d}"


def hm(s: float) -> str:
    s = int(round(s or 0))
    return f"{s // 3600}h {s % 3600 // 60:02d}m"


def select_videos(c, org_id: int, f: dict) -> list[dict]:
    q, args = "SELECT * FROM videos WHERE org_id=?", [org_id]
    if f.get("video_ids"):
        q += f" AND id IN ({','.join('?' * len(f['video_ids']))})"
        args += list(f["video_ids"])
    if f.get("date_from"):
        q += " AND recorded_on>=?"
        args.append(f["date_from"])
    if f.get("date_to"):
        q += " AND recorded_on<=?"
        args.append(f["date_to"])
    for key in ("site_id", "camera_id"):
        if f.get(key):
            q += f" AND {key}=?"
            args.append(int(f[key]))
    if f.get("shift"):
        q += " AND shift=?"
        args.append(f["shift"])
    return db.rows(c.execute(q + " ORDER BY recorded_on, id", args))


def aggregate(c, org_id: int, f: dict | None = None):
    """Per-person rows over the latest run of each selected video.
    Returns (rows, totals, zones, context) where context has used videos, notes and gate blockers."""
    f = f or {}
    rows: dict[str, dict] = {}
    zones_tot: dict[str, dict] = defaultdict(lambda: {s: 0.0 for s in STATES})
    used, notes, blockers = [], [], []
    for v in select_videos(c, org_id, f):
        v = store.video_view(c, v)
        if not v["run_id"]:
            if v["status"] in ("queued", "analysing", "uploaded", "slice"):
                blockers.append(f"{v['name']}: analysis not complete")
            continue
        if not v["annotated"]:
            blockers.append(f"{v['name']}: people & zones not annotated yet")
            if f.get("reviewed_only"):
                continue
        if v["pending"]:
            blockers.append(f"{v['name']}: {v['pending']} item(s) pending review")
            if f.get("reviewed_only"):
                continue
        used.append(v["id"])
        if v["partial"]:
            notes.append(f"{v['name']} ({v['recorded_on']}): partial footage — {v['gap_note'] or 'camera feed dropped'}")
        for p in store.run_people(c, v["run_id"], include_removed=False):
            name = p["worker_name"] or (p["label"] if not p["label"].startswith("Person ") else f"{v['name']} · {p['label']}")
            if f.get("worker") and f["worker"] != name:
                continue
            r = rows.setdefault(name, {"person": name, "employee_id": p.get("employee_id") or "", "videos": [],
                                       "dates": [], "flagged": 0, "pending": 0, "zones": defaultdict(lambda: {s: 0.0 for s in STATES}),
                                       **{f"{s}_s": 0.0 for s in STATES}})
            r["videos"].append(v["name"])
            r["dates"].append(v["recorded_on"])
            for s in p["segments"]:
                if f.get("zone") and (s["zone"] or "Outside zones") != f["zone"]:
                    continue
                st, d = s["effective_state"], s["end"] - s["start"]
                r[f"{st}_s"] += d
                zn = s["zone"] or "Outside zones"
                r["zones"][zn][st] += d
                zones_tot[zn][st] += d
                r["flagged"] += s["flagged"]
                r["pending"] += s["review_status"] == "pending"
    out = []
    for r in rows.values():
        t = {s: r[f"{s}_s"] for s in STATES}
        r["present_s"] = sum(t.values())
        if not r["present_s"]:
            continue
        r["efficiency"] = store.efficiency(t)
        r["videos"], r["dates"] = sorted(set(r["videos"])), sorted(set(r["dates"]))
        r["zones"] = {z: {k: round(v, 1) for k, v in d.items() if v} for z, d in r["zones"].items()}
        out.append(r)
    out.sort(key=lambda r: -r["working_s"])
    tot = {f"{s}_s": sum(r[f"{s}_s"] for r in out) for s in STATES}
    tot["present_s"] = sum(r["present_s"] for r in out)
    tot["people"] = len(out)
    tot["efficiency"] = store.efficiency({s: tot[f"{s}_s"] for s in STATES})
    tot["pending"] = sum(r["pending"] for r in out)
    zones = {z: {k: round(v, 1) for k, v in d.items() if v} for z, d in zones_tot.items()}
    return out, tot, zones, {"used": used, "notes": notes, "blockers": blockers}


def build(org_id: int, f: dict, user: dict) -> dict:
    with db.tx() as c:
        rows, totals, zones, ctx = aggregate(c, org_id, f)
        if ctx["blockers"]:
            raise GateError(f"{len(ctx['blockers'])} blocker(s): clear them to generate the report.", ctx["blockers"])
        if not rows:
            raise GateError("No analysed footage matches these filters.", [])
        dates = sorted({d for r in rows for d in r["dates"] if d})
        d0, d1 = f.get("date_from") or (dates[0] if dates else ""), f.get("date_to") or (dates[-1] if dates else "")
        title = f"Daily report — {d0}" if d0 == d1 else f"Report — {d0} to {d1}"
        f = {**f, "date_from": d0 or None, "date_to": d1 or None}
        rules = settings.get(c, org_id)["rules"]
        rid = c.execute(
            "INSERT INTO reports (org_id, title, report_date, date_to, filters, video_ids, rows, totals, zones, rules, "
            "notes, created_by, created_at) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (org_id, title, f.get("date_from"), f.get("date_to"), db.dumps(f), db.dumps(ctx["used"]), db.dumps(rows),
             db.dumps(totals), db.dumps(zones), db.dumps(rules), db.dumps(ctx["notes"]),
             user.get("email") or user.get("phone") or "system", db.now())).lastrowid
        db.audit(c, user, "report.generate", "report", rid, None, {"filters": f, "people": totals["people"]})
        events.emit(c, org_id, "report_ready", f"Report ready: {title}",
                    f"{totals['people']} people · {hm(totals['working_s'])} working · {totals['efficiency']}% efficiency",
                    "/reports")
        rep = db.row(c.execute("SELECT * FROM reports WHERE id=?", (rid,)).fetchone())
        on_ready = settings.get(c, org_id)["delivery"]["on_ready"]
    if on_ready:
        try:
            deliver(rep["id"], org_id, user)
        except Exception:
            log.exception("auto-delivery failed")
    return rep


# ---------- exports ----------

def _org(c, org_id):
    return db.row(c.execute("SELECT * FROM organizations WHERE id=?", (org_id,)).fetchone())


def to_csv(rep: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(["Person", "Employee ID", "Dates", "Footage"] + [f"{l} (h:mm:ss)" for _, l in COLS] +
               ["Working hours", "Efficiency %", "Flagged segments"])
    for r in rep["rows"]:
        w.writerow([r["person"], r["employee_id"], " ".join(r["dates"]), " | ".join(r["videos"])] +
                   [hms(r[f"{k}_s"]) for k, _ in COLS] + [f"{r['working_s'] / 3600:.2f}", r["efficiency"], r["flagged"]])
    t = rep["totals"]
    w.writerow(["TOTAL", "", "", ""] + [hms(t[f"{k}_s"]) for k, _ in COLS] + [f"{t['working_s'] / 3600:.2f}", t["efficiency"], ""])
    for n in rep["notes"]:
        w.writerow([f"Note: {n}"])
    return buf.getvalue().encode("utf-8-sig")


PAYROLL = {
    # Attendance-import layouts. Column names follow each product's standard attendance template;
    # confirm against the template downloaded from your account before the first import.
    "greytHR": (["Employee No", "Employee Name", "Attendance Date", "Status", "Worked Hours"],
                lambda r, d: [r["employee_id"], r["person"], d, "P" if r["working_s"] > 0 else "A", f"{r['working_s'] / 3600:.2f}"]),
    "Keka": (["Employee Number", "Employee Name", "Date", "Total Hours", "Effective Hours"],
             lambda r, d: [r["employee_id"], r["person"], d, f"{r['present_s'] / 3600:.2f}", f"{r['working_s'] / 3600:.2f}"]),
    "Zoho Payroll": (["Employee ID", "Employee Name", "Date", "Hours Worked"],
                     lambda r, d: [r["employee_id"], r["person"], d, f"{r['working_s'] / 3600:.2f}"]),
    "Tally": (["Employee Name", "Date", "Attendance Type", "Value"],
              lambda r, d: [r["person"], d, "Working Hours", f"{r['working_s'] / 3600:.2f}"]),
}


def to_payroll(rep: dict, fmt: str) -> bytes:
    head, fn = PAYROLL.get(fmt, PAYROLL["greytHR"])
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(head)
    for r in rep["rows"]:
        # Rows are already summed over the report's range, so write one line per person (never repeat the
        # same total on every date). Multi-day reports carry the range end date.
        d = r["dates"][0] if len(r["dates"]) == 1 else (rep.get("date_to") or rep.get("report_date") or "")
        w.writerow(fn(r, datetime.fromisoformat(d).strftime("%d-%m-%Y") if d else ""))
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(rep: dict, org: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    head_fill = PatternFill("solid", fgColor="1F2937")
    ws.append([f"{org['name']} — {rep['title']}"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"Generated {fmt_ts(rep['created_at'], org)} by {rep['created_by']} · times in {org['timezone']}"])
    ws.append([])
    header = ["Person", "Employee ID", "Dates"] + [f"{l} (h)" for _, l in COLS] + ["Efficiency %", "Flagged"]
    ws.append(header)
    for i in range(1, len(header) + 1):
        cell = ws.cell(row=4, column=i)
        cell.font, cell.fill, cell.alignment = Font(bold=True, color="FFFFFF"), head_fill, Alignment(wrap_text=True)
    for r in rep["rows"]:
        ws.append([r["person"], r["employee_id"], ", ".join(r["dates"])] +
                  [round(r[f"{k}_s"] / 3600, 3) for k, _ in COLS] + [r["efficiency"], r["flagged"]])
    t = rep["totals"]
    ws.append(["TOTAL", "", ""] + [round(t[f"{k}_s"] / 3600, 3) for k, _ in COLS] + [t["efficiency"], ""])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    for i, w in enumerate([28, 14, 22] + [12] * len(COLS) + [12, 10], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B5"

    zs = wb.create_sheet("Zone breakdown")
    zs.append(["Person", "Zone"] + [l + " (h)" for k, l in COLS if k != "present"])
    for r in rep["rows"]:
        for z, d in r["zones"].items():
            zs.append([r["person"], z] + [round(d.get(k, 0) / 3600, 3) for k, _ in COLS if k != "present"])
    for i, w in enumerate([28, 24, 12, 12, 16, 12, 12], 1):
        zs.column_dimensions[get_column_letter(i)].width = w

    pr = wb.create_sheet("Payroll")
    head, fn = PAYROLL["greytHR"]
    pr.append(head)
    for r in rep["rows"]:
        for d in r["dates"]:
            pr.append(fn(r, d))

    ns = wb.create_sheet("Notes & rules")
    ns.append(["Efficiency = Working ÷ (Present − Exclusion)"])
    for n in rep["notes"]:
        ns.append([n])
    ns.append([])
    ns.append(["Analysis rules used"])
    for k, v in (rep.get("rules") or {}).items():
        if k != "labels":
            ns.append([k, str(v)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _latin(s: str) -> str:
    return (str(s).replace("—", "-").replace("–", "-").replace("·", "|").replace("÷", "/").replace("₹", "Rs ")
            .replace("−", "-").encode("latin-1", "replace").decode("latin-1"))


def to_pdf(rep: dict, org: dict, footer: str = "") -> bytes:
    from fpdf import FPDF

    pdf = FPDF(orientation="L", format="A4")
    pdf.set_auto_page_break(True, 14)
    pdf.add_page()
    if org.get("logo") and (config.DATA_DIR / org["logo"]).exists():
        try:
            pdf.image(str(config.DATA_DIR / org["logo"]), x=250, y=8, h=14)
        except Exception:
            pass
    pdf.set_font("Helvetica", "B", 16)
    pdf.cell(0, 9, _latin(org["name"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 11)
    pdf.cell(0, 7, _latin(rep["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(110)
    pdf.set_font("Helvetica", "", 9)
    pdf.cell(0, 6, _latin(f"Generated {fmt_ts(rep['created_at'], org)} by {rep['created_by']} | times in {org['timezone']} | "
                          f"all flagged segments reviewed"), new_x="LMARGIN", new_y="NEXT")
    pdf.set_text_color(0)
    t = rep["totals"]
    pdf.ln(2)
    pdf.set_font("Helvetica", "B", 10)
    pdf.cell(0, 7, _latin(f"{t['people']} people   Present {hm(t['present_s'])}   Working {hm(t['working_s'])}   "
                          f"Idle {hm(t['idle_s'])}   Phone {hm(t['phone_s'])}   Break {hm(t['break_s'])}   "
                          f"Efficiency {t['efficiency']}%"), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    widths = [62, 26] + [24] * len(COLS) + [22, 18]
    head = ["Person", "Emp. ID"] + [l for _, l in COLS] + ["Effic.", "Flags"]
    pdf.set_font("Helvetica", "B", 8.5)
    pdf.set_fill_color(31, 41, 55)
    pdf.set_text_color(255)
    for w, h in zip(widths, head):
        pdf.cell(w, 8, _latin(h), border=0, fill=True)
    pdf.ln()
    pdf.set_text_color(0)
    pdf.set_font("Helvetica", "", 8.5)
    for i, r in enumerate(rep["rows"]):
        pdf.set_fill_color(*((245, 246, 248) if i % 2 else (255, 255, 255)))
        vals = [r["person"][:38], r["employee_id"]] + [hm(r[f"{k}_s"]) for k, _ in COLS] + [f"{r['efficiency']}%", str(r["flagged"])]
        for w, val in zip(widths, vals):
            pdf.cell(w, 7, _latin(val), fill=True)
        pdf.ln()
    if any(r["zones"] for r in rep["rows"]):
        pdf.ln(4)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 7, "Zone breakdown (hours)", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8.5)
        for r in rep["rows"]:
            for z, d in r["zones"].items():
                parts = ", ".join(f"{l} {d[k] / 3600:.2f}" for k, l in COLS if d.get(k))
                pdf.cell(0, 6, _latin(f"{r['person']} | {z}: {parts}"), new_x="LMARGIN", new_y="NEXT")
    if rep["notes"]:
        pdf.ln(3)
        pdf.set_font("Helvetica", "B", 10)
        pdf.cell(0, 7, "Notes", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("Helvetica", "", 8.5)
        for n in rep["notes"]:
            pdf.multi_cell(0, 5, _latin(n), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    pdf.set_text_color(120)
    pdf.set_font("Helvetica", "", 8)
    pdf.multi_cell(0, 4.5, _latin("Efficiency = Working / (Present - Exclusion). " + (footer or "")),
                   new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


def fmt_ts(ts: float, org: dict) -> str:
    return datetime.fromtimestamp(ts, ZoneInfo(org.get("timezone") or "Asia/Kolkata")).strftime("%d %b %Y, %I:%M %p")


def export(rep: dict, fmt: str, payroll: str | None = None) -> tuple[bytes, str, str]:
    with db.tx() as c:
        org = _org(c, rep["org_id"])
        footer = settings.get(c, rep["org_id"])["delivery"]["pdf_footer"]
    base = f"workwatch_report_{rep['id']}_{rep.get('report_date') or 'all'}"
    if fmt == "xlsx":
        return to_xlsx(rep, org), f"{base}.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    if fmt == "pdf":
        return to_pdf(rep, org, footer), f"{base}.pdf", "application/pdf"
    if fmt == "payroll":
        p = payroll or "greytHR"
        return to_payroll(rep, p), f"{base}_{p.replace(' ', '_').lower()}.csv", "text/csv"
    return to_csv(rep), f"{base}.csv", "text/csv"


# ---------- delivery ----------

def deliver(report_id: int, org_id: int, user: dict | None, channels=("email", "whatsapp")) -> list[dict]:
    with db.tx() as c:
        rep = db.row(c.execute("SELECT * FROM reports WHERE id=? AND org_id=?", (report_id, org_id)).fetchone())
        s = settings.get(c, org_id)["delivery"]
        managers = [r["email"] for r in c.execute(
            "SELECT email FROM users WHERE org_id=? AND role='manager' AND status='active' AND email IS NOT NULL", (org_id,))]
    t = rep["totals"]
    summary = (f"{rep['title']}\n{t['people']} people | Working {hm(t['working_s'])} | Idle {hm(t['idle_s'])} | "
               f"Phone {hm(t['phone_s'])} | Break {hm(t['break_s'])} | Efficiency {t['efficiency']}%")
    if rep["notes"]:
        summary += "\nNotes: " + "; ".join(rep["notes"])
    results = []

    def log_delivery(channel, recipient, status, error=None, pid=None):
        with db.tx() as c:
            c.execute("INSERT INTO deliveries (org_id, report_id, channel, recipient, status, error, provider_id, created_at) "
                      "VALUES (?,?,?,?,?,?,?,?)", (org_id, report_id, channel, recipient, status, error, pid, db.now()))
        results.append({"channel": channel, "recipient": recipient, "status": status, "error": error})

    if "email" in channels:
        to = list(dict.fromkeys(s["email_recipients"] or managers))
        if not delivery.email_configured():
            log_delivery("email", ", ".join(to) or "—", "not_configured", "SMTP is not configured")
        elif to:
            atts = [(name, data, mime) for data, name, mime in (export(rep, f) for f in s["formats"])]
            rows_html = "".join(f"<tr><td>{r['person']}</td><td>{hm(r['working_s'])}</td><td>{hm(r['idle_s'])}</td>"
                                f"<td>{hm(r['phone_s'])}</td><td>{hm(r['break_s'])}</td><td>{r['efficiency']}%</td></tr>"
                                for r in rep["rows"])
            html = (f"<h2>{rep['title']}</h2><p>{summary.splitlines()[1]}</p><table border=1 cellpadding=6 "
                    f"style='border-collapse:collapse'><tr><th>Person</th><th>Working</th><th>Idle</th><th>Phone</th>"
                    f"<th>Break</th><th>Efficiency</th></tr>{rows_html}</table>")
            try:
                delivery.send_email(to, f"WorkWatch: {rep['title']}", summary, html, atts)
                log_delivery("email", ", ".join(to), "sent")
            except Exception as e:
                log_delivery("email", ", ".join(to), "failed", str(e)[:300])
    if "whatsapp" in channels:
        rcps = [r for r in s["whatsapp_recipients"] if r.get("opted_in")]
        if not delivery.whatsapp_configured():
            log_delivery("whatsapp", ", ".join(r["number"] for r in rcps) or "—", "not_configured",
                         "WhatsApp Business API is not configured")
        for r in rcps if delivery.whatsapp_configured() else []:
            link = f"{config.PUBLIC_URL}/r/{delivery.sign({'r': report_id, 'f': 'pdf'})}"
            try:
                pid = delivery.send_whatsapp(r["number"], summary, document_url=link, filename=f"report_{report_id}.pdf")
                log_delivery("whatsapp", r["number"], "sent", pid=pid)
            except Exception as e:
                log_delivery("whatsapp", r["number"], "failed", str(e)[:300])
    if any(x["status"] == "sent" for x in results):
        with db.tx() as c:
            c.execute("UPDATE reports SET status='sent' WHERE id=?", (report_id,))
            ids = rep["video_ids"]
            if ids:
                c.execute(f"UPDATE videos SET sent=1 WHERE id IN ({','.join('?' * len(ids))})", ids)
            if user:
                db.audit(c, user, "report.deliver", "report", report_id, None, results)
    return results


def send_summary(org_id: int, text: str, kind: str) -> None:
    with db.tx() as c:
        s = settings.get(c, org_id)["delivery"]
        if not s.get(kind, True):
            return
    for r in s["whatsapp_recipients"]:
        if r.get("opted_in") and delivery.whatsapp_configured():
            try:
                delivery.send_whatsapp(r["number"], text)
            except Exception:
                log.exception("weekly WhatsApp summary failed")
    if s["email_recipients"] and delivery.email_configured():
        try:
            delivery.send_email(s["email_recipients"], f"WorkWatch {kind} summary", text)
        except Exception:
            log.exception("weekly email summary failed")
